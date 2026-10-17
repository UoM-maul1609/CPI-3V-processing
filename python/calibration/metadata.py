"""Read calibration truth tables without altering measured CPI quantities."""
from __future__ import annotations

import csv
import math
from pathlib import Path


def _normalise(name):
    return str(name or '').strip().lower().replace(' ', '_').replace('-', '_')


def read_truth_table(filename):
    """Rows keyed by particle_id or source_file (+ optional roi_index).

    Expected columns: source_file or particle_id; actual_length_um,
    actual_width_um, object_plane_um. An optional instrument column is
    CPI1.0 or CPI3v. Coordinates are stored *as supplied*, not flipped.
    """
    path = Path(filename)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
    elif path.suffix.lower() == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('Install openpyxl for Excel import: pip install openpyxl') from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            values = wb.active.iter_rows(values_only=True)
            header = next(values)
            rows = [dict(zip(header, values_row)) for values_row in values]
        finally:
            wb.close()
    else:
        raise ValueError('Use a .csv or .xlsx spreadsheet')
    result = []
    for number, row in enumerate(rows, 2):
        entry = {_normalise(k): v for k, v in row.items()}
        if not any(v is not None and str(v).strip() for v in entry.values()):
            continue
        for alias, canonical in [('length_um','actual_length_um'), ('width_um','actual_width_um'),
                                  ('z_um','object_plane_um'), ('plane_position_um','object_plane_um')]:
            if canonical not in entry and alias in entry:
                entry[canonical] = entry[alias]
        if not (entry.get('particle_id') or entry.get('source_file')):
            raise ValueError(f'Row {number}: specify source_file or particle_id')
        for key in ('actual_length_um', 'actual_width_um', 'object_plane_um'):
            value = entry.get(key)
            if value is None or str(value).strip() == '':
                continue
            try:
                entry[key] = float(value)
            except (ValueError, TypeError) as exc:
                raise ValueError(f'Row {number}: invalid {key}: {value}') from exc
            if not math.isfinite(entry[key]):
                raise ValueError(f'Row {number}: {key} must be finite')
        result.append(entry)
    return result


def attach_truth(records, rows):
    """Apply exact particle / source matches, rejecting ambiguous matches."""
    matched = 0
    unmatched = []
    for row in rows:
        key = str(row.get('particle_id') or '').strip()
        source = str(row.get('source_file') or '').strip()
        roi = row.get('roi_index')
        candidates = [r for r in records if (r.particle_id == key if key else
            (r.source_file == source or Path(r.source_file).name == source) and
            (roi is None or str(roi).strip() == '' or r.roi_index == int(roi)))]
        if not candidates:
            unmatched.append(key or source)
            continue
        # Rows keyed by a source without ROI deliberately apply to all its particles.
        if len({r.source_file for r in candidates}) > 1:
            raise ValueError(f'Ambiguous source {source!r}; use full path or particle_id')
        for record in candidates:
            for field in ('actual_length_um', 'actual_width_um', 'object_plane_um'):
                if row.get(field) is not None and str(row[field]).strip() != '':
                    setattr(record, field, float(row[field]))
            instrument = str(row.get('instrument') or '').strip()
            if instrument:
                if instrument.lower() not in ('cpi1.0', 'cpi3v'):
                    raise ValueError(f'Unknown instrument: {instrument}')
                record.instrument = 'CPI1.0' if instrument.lower() == 'cpi1.0' else 'CPI3v'
            matched += 1
    return matched, unmatched
